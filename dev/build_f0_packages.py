"""Render the F0 decision packages (plan «expediente completo») from the JSON of the drafting workflow:
docs/decisiones-f0.md (in Git), one Excel per person docs/decisiones-f0-<owner>.xlsx (outside Git: the repository
ignores spreadsheets) with yellow answer columns, and the register section of docs/preguntas-abiertas.md.
    python dev/build_f0_packages.py <packages.json>"""
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
packages = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
ORDER = {"Ángel": 0, "Cristina": 1, "Belal": 2}
SLUG = {"Ángel": "angel", "Cristina": "cristina", "Belal": "belal"}
STAGE_ORDER = ["0", "1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11", "P"]


def stage_key(stage):
    head = stage.split(" ")[0].strip(".·:")
    return STAGE_ORDER.index(head) if head in STAGE_ORDER else len(STAGE_ORDER)


def day_key(deadline):
    try:
        d, m, y = deadline.split("/")
        return (int(y), int(m), int(d))
    except ValueError:
        return (9999, 0, 0)


packages = sorted(packages, key=lambda p: ORDER.get(p["owner"], 9))
for pkg in packages:
    pkg["questions"].sort(key=lambda q: (stage_key(q["stage"]), day_key(q["deadline"]), q["id"]))

# ---- markdown -------------------------------------------------------------------------------------------------------
md = ["# Decisiones F0 del plan «expediente completo» (28/10/2026)", "",
      "Preguntas que bloquean las fases del plan (`docs/plan-expediente-completo.md`), agrupadas por persona y por etapa "
      "del expediente. **Nada de esto está decidido**: «Hoy en Odoo 17» es material para la pregunta, no la regla, y "
      "«Mientras tanto» es lo que hace Odoo 20 hasta la respuesta (estado explícito, nunca una suposición). Cada "
      "respuesta se registra con nombre y fecha en `docs/preguntas-abiertas.md`.", ""]
md.append("| Persona | Preguntas | Prioridad alta | Primera fecha límite |")
md.append("|---|---|---|---|")
for pkg in packages:
    qs = pkg["questions"]
    first = min((q["deadline"] for q in qs), key=day_key) if qs else "—"
    md.append(f"| {pkg['owner']} | {len(qs)} | {sum(q['priority'] == 'alta' for q in qs)} | {first} |")
md.append("")
for pkg in packages:
    md.append(f"## {pkg['owner']}")
    current = None
    for q in pkg["questions"]:
        if q["stage"] != current:
            current = q["stage"]
            md.append(f"### Etapa {current}")
        reused = " (ya abierta)" if q["existing_id_reused"] else ""
        md.append(f"**{q['id']}{reused} · {q['priority']} · antes del {q['deadline']}** — {q['question_es']}")
        md.append("")
        md.append(f"- Por qué: {q['why_es']}")
        if q["options_es"]:
            md.append("- Opciones: " + " / ".join(q["options_es"]))
        if q["odoo17_es"]:
            md.append(f"- Hoy en Odoo 17: {q['odoo17_es']}")
        md.append(f"- Mientras tanto: {q['meanwhile_es']}")
        md.append(f"- Bloquea: {', '.join(q['blocks_phases'])}" + (f" · con {q['with_also']}" if q["with_also"] else ""))
        md.append("")
(ROOT / "docs" / "decisiones-f0.md").write_text("\n".join(md) + "\n", encoding="utf-8", newline="\n")

# ---- one Excel per person -----------------------------------------------------------------------------------------------
FONT, NAVY, INPUT = "Arial", "14324A", "FFF9C4"
thin = Side(style="thin", color="D5DBE1")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
HEAD = ["Id", "Etapa del expediente", "Pregunta", "Por qué importa / qué bloquea", "Opciones", "Hoy en Odoo 17",
        "Mientras tanto en Odoo 20", "Bloquea", "Fecha límite", "Prioridad", "Con", "Decisión", "Comentarios"]
WIDTHS = (9, 18, 46, 40, 34, 32, 32, 10, 11, 9, 11, 34, 34)
for pkg in packages:
    wb = Workbook()
    ws = wb.active
    ws.title = "Decisiones"
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"BOSS Odoo 20 · Decisiones pendientes de {pkg['owner']} · plan «expediente completo» (28/10/2026)"
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    ws["A2"] = ("Escriba su decisión y sus comentarios en las dos columnas amarillas y devuelva el archivo a Belal. "
                "«Hoy en Odoo 17» es solo información; nada está decidido hasta su respuesta.")
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="6B7785")
    for col, (text, width) in enumerate(zip(HEAD, WIDTHS), start=1):
        cell = ws.cell(row=4, column=col, value=text)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY if col < 12 else "8A7A00")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = width
    for r, q in enumerate(pkg["questions"], start=5):
        values = [q["id"], q["stage"], q["question_es"], q["why_es"], "\n".join(f"• {o}" for o in q["options_es"]),
                  q["odoo17_es"] or "—", q["meanwhile_es"], ", ".join(q["blocks_phases"]), q["deadline"],
                  q["priority"], q["with_also"] or "—", None, None]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = Font(name=FONT, size=10, bold=(col in (1, 3)))
            cell.alignment = wrap
            cell.border = border
            if col >= 12:
                cell.fill = PatternFill("solid", fgColor=INPUT)
            if col == 10 and value == "alta":
                cell.font = Font(name=FONT, size=10, bold=True, color="B42318")
    last = 4 + len(pkg["questions"])
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:M{last}"
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    out = ROOT / "docs" / f"decisiones-f0-{SLUG.get(pkg['owner'], pkg['owner'].lower())}.xlsx"
    wb.save(out)

# ---- register in docs/preguntas-abiertas.md ---------------------------------------------------------------------------
reg = ROOT / "docs" / "preguntas-abiertas.md"
text = reg.read_text(encoding="utf-8")
marker = "## Decisiones F0 del plan «expediente completo» (28/10/2026)"
if marker in text:
    text = text[:text.index(marker)].rstrip() + "\n"
rows = [marker, "", "Registro de las preguntas enviadas en F0 (detalle y opciones: `docs/decisiones-f0.md`). Estado: "
        "**abierta** hasta que la persona responda; la respuesta se anota aquí con nombre y fecha.", "",
        "| Id | Responsable | Etapa | Pregunta | Fecha límite | Mientras tanto | Estado |", "|---|---|---|---|---|---|---|"]
for pkg in packages:
    for q in pkg["questions"]:
        owner = pkg["owner"] + (f" + {q['with_also']}" if q["with_also"] else "")
        question = q["question_es"].replace("|", "/")
        meanwhile = q["meanwhile_es"].replace("|", "/")
        rows.append(f"| {q['id']} | {owner} | {q['stage']} | {question} | {q['deadline']} | {meanwhile} | abierta |")
reg.write_text(text.rstrip() + "\n\n" + "\n".join(rows) + "\n", encoding="utf-8", newline="\n")
print({p["owner"]: len(p["questions"]) for p in packages},
      "reused ids:", sum(q["existing_id_reused"] for p in packages for q in p["questions"]),
      "priorities:", dict(Counter(q["priority"] for p in packages for q in p["questions"])))
