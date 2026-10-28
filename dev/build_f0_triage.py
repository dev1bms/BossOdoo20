"""Render the triaged F0 decisions (Belal, 28/10/2026: ask the owners only what matters).
Inputs: .research/f0_packages.json (the questions) and .research/f0_triage_<chunk>.json (verdicts: answered, default, keep,
duplicate — each with its answer and sources). Outputs:
  docs/decisiones-f0.md      meeting list per person, adopted defaults, questions already answered (all with sources)
  docs/decisiones-f0-<owner>.xlsx  only the meeting questions, with answer columns (outside Git)
  docs/preguntas-abiertas.md  register section with the status of every question
    python dev/build_f0_triage.py"""
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
RES = ROOT / ".research"
packages = json.loads((RES / "f0_packages.json").read_text(encoding="utf-8"))
questions = {(p["owner"], q["id"]): q for p in packages for q in p["questions"]}
verdicts = {}
for chunk, owner in (("angel", "Ángel"), ("cristina1", "Cristina"), ("cristina2", "Cristina"), ("belal", "Belal")):
    for v in json.loads((RES / f"f0_triage_{chunk}.json").read_text(encoding="utf-8")):
        verdicts[(owner, v["id"])] = v
SLUG = {"Ángel": "angel", "Cristina": "cristina", "Belal": "belal"}
LABEL = {"answered": "respondida", "default": "supuesto adoptado", "keep": "reunión", "duplicate": "duplicada"}


def cell(text):
    return (text or "").replace("|", "/").replace("\n", " ")


rows = []
for (owner, qid), q in questions.items():
    v = verdicts.get((owner, qid), {"verdict": "keep", "answer_es": q["question_es"], "source": "sin revisar",
                                    "why_keep_es": q["why_es"], "confidence": "baja", "duplicate_of": ""})
    rows.append((owner, qid, q, v))
count = Counter((owner, v["verdict"]) for owner, _qid, _q, v in rows)

md = ["# Decisiones F0 del plan «expediente completo» — revisadas (28/10/2026)", "",
      "Belal no puede consultar a Cristina ni a Ángel por cada duda. Por eso cada pregunta se ha buscado primero en la "
      "documentación del proyecto, el proyecto Odoo 19, los Excel de SharePoint («NEW ODOO»), Odoo 17 de producción y, "
      "si no había respuesta, en la práctica del sector y la documentación de Odoo. Resultado en tres grupos:", "",
      "1. **Para la reunión**: esenciales, sin respuesta en ninguna fuente, y un supuesto equivocado causaría daño "
      "legal, económico o irreversible.",
      "2. **Supuestos adoptados**: no esenciales; se aplica la mejor práctica documentada, es reversible y queda a la vista "
      "del responsable para corregirla si no está de acuerdo.",
      "3. **Ya respondidas**: la respuesta existe en una fuente (citada).", ""]
md.append("| Persona | Para la reunión | Supuestos adoptados | Ya respondidas | Duplicadas |")
md.append("|---|---|---|---|---|")
for owner in ("Ángel", "Cristina", "Belal"):
    md.append(f"| {owner} | {count[(owner, 'keep')]} | {count[(owner, 'default')]} | {count[(owner, 'answered')]} | "
              f"{count[(owner, 'duplicate')]} |")
md.append("")
md.append("## 1. Para la reunión")
for owner in ("Ángel", "Cristina", "Belal"):
    keep = [(qid, q, v) for o, qid, q, v in rows if o == owner and v["verdict"] == "keep"]
    md.append(f"### {owner} ({len(keep)})")
    for qid, q, v in keep:
        md.append(f"**{qid}** — {v['answer_es']}")
        md.append(f"- Por qué hay que preguntarlo: {v.get('why_keep_es') or q['why_es']}")
        if q["options_es"]:
            md.append("- Opciones: " + " / ".join(q["options_es"]))
        md.append(f"- Mientras tanto: {q['meanwhile_es']}")
        md.append(f"- Buscado en: {v.get('source') or '—'}")
        md.append("")
md.append("## 2. Supuestos adoptados (reversibles)")
md.append("| Persona | Id | Pregunta | Supuesto adoptado | Fuente | Confianza |")
md.append("|---|---|---|---|---|---|")
for owner, qid, q, v in rows:
    if v["verdict"] == "default":
        md.append(f"| {owner} | {qid} | {cell(q['question_es'])} | {cell(v['answer_es'])} | {cell(v.get('source'))} | "
                  f"{v.get('confidence', '')} |")
md.append("")
md.append("## 3. Ya respondidas")
md.append("| Persona | Id | Pregunta | Respuesta | Fuente |")
md.append("|---|---|---|---|---|")
for owner, qid, q, v in rows:
    if v["verdict"] == "answered":
        md.append(f"| {owner} | {qid} | {cell(q['question_es'])} | {cell(v['answer_es'])} | {cell(v.get('source'))} |")
md.append("")
dups = [(owner, qid, v) for owner, qid, _q, v in rows if v["verdict"] == "duplicate"]
if dups:
    md.append("Duplicadas (se responden con otra pregunta): " + ", ".join(f"{qid} → {v.get('duplicate_of')}" for _o, qid, v in dups))
(ROOT / "docs" / "decisiones-f0.md").write_text("\n".join(md) + "\n", encoding="utf-8", newline="\n")

# ---- Excel: only the meeting questions ---------------------------------------------------------------------------------
FONT, NAVY, INPUT = "Arial", "14324A", "FFF9C4"
thin = Side(style="thin", color="D5DBE1")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
HEAD = ["Id", "Pregunta", "Por qué hay que preguntarlo", "Opciones", "Mientras tanto en Odoo 20", "Bloquea",
        "Fecha límite", "Decisión", "Comentarios"]
WIDTHS = (9, 52, 44, 38, 36, 10, 11, 34, 34)
for owner in ("Ángel", "Cristina", "Belal"):
    keep = [(qid, q, v) for o, qid, q, v in rows if o == owner and v["verdict"] == "keep"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Para la reunión"
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"BOSS Odoo 20 · Preguntas para {owner} (28/10/2026)"
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    ws["A2"] = ("Solo lo esencial que no está respondido en ninguna fuente. Lo demás se ha resuelto con la documentación "
                "o con un supuesto reversible (docs/decisiones-f0.md). Escriba la decisión en las columnas amarillas.")
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="6B7785")
    for col, (text, width) in enumerate(zip(HEAD, WIDTHS), start=1):
        c = ws.cell(row=4, column=col, value=text)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY if col < 8 else "8A7A00")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = border
        ws.column_dimensions[c.column_letter].width = width
    for r, (qid, q, v) in enumerate(keep, start=5):
        values = [qid, v["answer_es"], v.get("why_keep_es") or q["why_es"], "\n".join(f"• {o}" for o in q["options_es"]),
                  q["meanwhile_es"], ", ".join(q["blocks_phases"]), q["deadline"], None, None]
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=value)
            c.font = Font(name=FONT, size=10, bold=(col in (1, 2)))
            c.alignment = wrap
            c.border = border
            if col >= 8:
                c.fill = PatternFill("solid", fgColor=INPUT)
    ws.freeze_panes = "C5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(ROOT / "docs" / f"decisiones-f0-{SLUG[owner]}.xlsx")

# ---- register in docs/preguntas-abiertas.md -------------------------------------------------------------------------------
reg = ROOT / "docs" / "preguntas-abiertas.md"
text = reg.read_text(encoding="utf-8")
marker = "## Decisiones F0 del plan «expediente completo» (28/10/2026)"
if marker in text:
    text = text[:text.index(marker)].rstrip() + "\n"
lines = [marker, "", "Estado de cada pregunta tras la revisión del 28/10 (detalle y fuentes: `docs/decisiones-f0.md`). «reunión» = "
         "abierta para la reunión; «supuesto adoptado» = se aplica la mejor práctica, reversible; «respondida» = ya estaba "
         "decidida en una fuente.", "", "| Id | Responsable | Pregunta | Estado | Respuesta o supuesto |", "|---|---|---|---|---|"]
for owner, qid, q, v in rows:
    answer = "" if v["verdict"] == "keep" else cell(v["answer_es"])
    lines.append(f"| {qid} | {owner} | {cell(q['question_es'])} | {LABEL.get(v['verdict'], v['verdict'])} | {answer} |")
reg.write_text(text.rstrip() + "\n\n" + "\n".join(lines) + "\n", encoding="utf-8", newline="\n")
print({o: {k: count[(o, k)] for k in ("keep", "default", "answered", "duplicate")} for o in ("Ángel", "Cristina", "Belal")})
