"""Build docs/angel-meetings-status.xlsx: one sheet, every point of the two meetings with Ángel and its status
(same data as the HTML: dev/build_meeting_status_html.py).
    python dev/build_meeting_status_xlsx.py"""
import importlib.util
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("status_html", HERE / "build_meeting_status_html.py")
src = importlib.util.module_from_spec(spec)
spec.loader.exec_module(src)  # also rewrites the HTML from the same rows
ROWS = src.ROWS
STATUSES = [(src.CLOSED, "1E7F4F", "E3F4EA"), (src.SCREEN, "2563A8", "E4EEFA"), (src.PARTIAL, "A15C07", "FDF0DC"),
            (src.DECISION, "7A3FA0", "F1E7F8"), (src.OPEN, "B42318", "FDE8E7")]

FONT = "Arial"
thin = Side(style="thin", color="D5DBE1")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
ws = wb.active
ws.title = "Status"
ws.sheet_view.showGridLines = False

ws["A1"] = "BOSS Odoo 20 · Meetings with Ángel (27/10/2026) · Status of every point"
ws["A1"].font = Font(name=FONT, size=14, bold=True, color="14324A")
ws["A2"] = "Status on the boss20 server, 28/10/2026. Meeting 1 = review checklist (42 points); Meeting 2 = seven recovered segments (23 points)."
ws["A2"].font = Font(name=FONT, size=9, italic=True, color="6B7785")

# summary (formulas count the table below)
first, last = 13, 13 + len(ROWS) - 1
status_range = f"$D${first}:$D${last}"
ws["A4"], ws["C4"], ws["D4"] = "Status", "Points", "Share"
for col in ("A4", "C4", "D4"):
    ws[col].font = Font(name=FONT, bold=True, color="FFFFFF")
    ws[col].fill = PatternFill("solid", fgColor="14324A")
ws.merge_cells("A4:B4")
for i, (status, ink, bg) in enumerate(STATUSES, start=5):
    ws.merge_cells(f"A{i}:B{i}")
    ws[f"A{i}"] = status
    ws[f"A{i}"].font = Font(name=FONT, bold=True, color=ink)
    ws[f"A{i}"].fill = PatternFill("solid", fgColor=bg)
    ws[f"C{i}"] = f'=COUNTIF({status_range},A{i})'
    ws[f"D{i}"] = f"=IF($C$10=0,0,C{i}/$C$10)"
ws.merge_cells("A10:B10")
ws["A10"], ws["C10"], ws["D10"] = "Total", f"=SUM(C5:C9)", "=SUM(D5:D9)"
ws.merge_cells("A11:B11")
ws["A11"], ws["D11"] = "Technical completion (closed + implemented)", "=IF($C$10=0,0,(C5+C6)/$C$10)"
for r in range(5, 12):
    for c in "ABCD":
        cell = ws[f"{c}{r}"]
        cell.border = border
        if c != "A" or r >= 10:
            cell.font = Font(name=FONT, bold=r >= 10)
    ws[f"D{r}"].number_format = "0%"
    ws[f"C{r}"].alignment = Alignment(horizontal="center")
ws["D11"].font = Font(name=FONT, bold=True, size=12, color="1E7F4F")

# table
headers = ["#", "Meeting", "Point raised", "Status", "Pending · owner"]
for col, text in enumerate(headers, start=1):
    cell = ws.cell(row=12, column=col, value=text)
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="14324A")
    cell.alignment = Alignment(vertical="center")
    cell.border = border
for r, (pid, meeting, topic, status, _done, pending) in enumerate(ROWS, start=first):
    values = [pid, meeting, topic, status, pending or "—"]
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=r, column=col, value=value)
        cell.font = Font(name=FONT, size=10, bold=(col == 1))
        cell.alignment = wrap
        cell.border = border
for status, ink, bg in STATUSES:
    ws.conditional_formatting.add(f"D{first}:D{last}", CellIsRule(
        operator="equal", formula=[f'"{status}"'], font=Font(name=FONT, bold=True, color=ink),
        fill=PatternFill("solid", fgColor=bg)))
for col, width in zip("ABCDE", (7, 11, 52, 30, 46)):
    ws.column_dimensions[col].width = width
ws.freeze_panes = "A13"
ws.auto_filter.ref = f"A12:E{last}"
ws.print_title_rows = "12:12"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.cell(row=last + 2, column=1,
        value="Closed = built, tested on boss20 and seen on screen · Implemented = built and tested, screen check pending · "
              "Partial = business condition open · Decision needed = Ángel or Cristina · Open = not started. "
              "Source: docs/matriz-reunion-2026-10-27.md").font = Font(name=FONT, size=8, italic=True, color="6B7785")

out = HERE.parent / "docs" / "angel-meetings-status.xlsx"
wb.save(out)
print(out.name, len(ROWS), "rows")
